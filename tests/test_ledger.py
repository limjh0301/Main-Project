import pytest

from extractor.ledger import (
    HEADER,
    MEMBER_SHEET_NAME,
    LedgerFormatError,
    append_rows,
    create_ledger,
    enrich_member_fields,
    open_ledger,
)

ROW = {
    "committee": "기획재정위원회",
    "request_date": "2024-01-01",
    "deadline": "2024.02.01",
    "doc_no": "0001234567",
    "member": "을지문덕",
    "party": "칼퇴원한당",
    "district": "조퇴도 원할시갑",
    "content": "1. 최근 5년간 직원 칼퇴 내역\n2. 유연근무제 운영 현황",
    "requester": "박삼관",
    "email": "admin@yogu.work",
}


def test_create_ledger_header():
    wb = create_ledger()
    ws = wb.active
    assert [ws.cell(row=1, column=i + 1).value for i in range(len(HEADER))] == HEADER
    assert HEADER[:2] == ["요구서번호", "소관위원회"]


def test_create_ledger_has_member_sheet():
    wb = create_ledger()
    assert MEMBER_SHEET_NAME in wb.sheetnames
    ws = wb[MEMBER_SHEET_NAME]
    assert ws.cell(row=1, column=3).value == "의원명"
    assert ws.max_row > 200  # 제22대 의원 명단 수록


def test_append_one_row_per_request(tmp_path):
    wb = create_ledger()
    assert append_rows(wb, [ROW, ROW]) == 2

    path = tmp_path / "ledger.xlsx"
    wb.save(path)
    wb2 = open_ledger(str(path))
    append_rows(wb2, [ROW])
    ws = wb2[wb2.sheetnames[0]]
    assert ws.max_row == 4  # 헤더 + 3행 (요구서 3건)
    assert ws.cell(row=2, column=1).value == "0001234567"
    assert ws.cell(row=2, column=8).value.startswith("1. 최근 5년간")
    # 수동 관리 항목(담당부서·제출여부·비고)은 비워 둔다
    assert ws.cell(row=2, column=11).value in ("", None)


def test_lookup_formula_links_member_sheet():
    wb = create_ledger()
    append_rows(wb, [ROW])
    ws = wb.active
    # 소관위원회(B)·정당(F)·지역구(G)는 의원명(E) 기준 VLOOKUP + 파싱값 폴백
    committee = ws.cell(row=2, column=2).value
    assert committee.startswith("=IFERROR(VLOOKUP($E2,")
    assert MEMBER_SHEET_NAME in committee
    assert committee.endswith(',"기획재정위원회")')
    assert ws.cell(row=2, column=6).value.endswith(',"칼퇴원한당")')
    assert ws.cell(row=2, column=7).value.endswith(',"조퇴도 원할시갑")')


def test_enrich_member_fields():
    data = {"member": "황희", "party": "", "committee": "", "district": ""}
    enrich_member_fields(data)
    assert data["party"] == "더불어민주당"
    assert data["committee"] == "국토교통위원회"
    assert data["district"] == "서울 양천구갑"

    unknown = {"member": "을지문덕", "party": "칼퇴원한당", "committee": "", "district": ""}
    enrich_member_fields(unknown)
    assert unknown["party"] == "칼퇴원한당"  # 명단에 없으면 파싱값 유지


def test_open_ledger_adds_missing_member_sheet(tmp_path):
    wb = create_ledger()
    del wb[MEMBER_SHEET_NAME]
    path = tmp_path / "old.xlsx"
    wb.save(path)
    wb2 = open_ledger(str(path))
    assert MEMBER_SHEET_NAME in wb2.sheetnames


def test_open_ledger_rejects_wrong_header(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.append(["잘못된", "헤더"])
    path = tmp_path / "bad.xlsx"
    wb.save(path)
    with pytest.raises(LedgerFormatError):
        open_ledger(str(path))
