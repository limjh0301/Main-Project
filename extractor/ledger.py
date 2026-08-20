"""엑셀 관리대장 생성·누적 모듈 (openpyxl).

관리대장 열 구조는 시스템이 자동 생성하며, 임의로 열을 바꾸면 기능 오류가
발생할 수 있으므로 기존 파일을 열 때 헤더를 검증한다.
요구서 1건 = 1행으로 기록한다.

'국회의원 현황' 시트(제22대 의원 명단)를 함께 넣고, 소관위원회·정당·지역구는
요구의원 이름을 키로 한 VLOOKUP 수식으로 연동한다. 명단에 없는 요구자(기관 등)는
요구서에서 파싱한 값으로 대신 채운다.
"""

import csv
import os

SHEET_NAME = "관리대장"
MEMBER_SHEET_NAME = "국회의원 현황"
_MEMBER_CSV = os.path.join(os.path.dirname(__file__), "data", "assembly_members.csv")

# (헤더, 너비, 행 데이터 키) — 키가 None이면 수동 관리 항목(빈 값으로 둠)
COLUMNS = [
    ("요구서번호", 14, "doc_no"),
    ("소관위원회", 18, "committee"),
    ("요구일자", 12, "request_date"),
    ("제출기한", 12, "deadline"),
    ("요구의원/기관", 14, "member"),
    ("정당", 14, "party"),
    ("지역구", 16, "district"),
    ("자료 요구내용", 55, "content"),
    ("요구자", 10, "requester"),
    ("요구자 이메일", 24, "email"),
    ("담당 부서", 18, None),
    ("제출여부", 10, None),
    ("비고/메모", 20, None),
]

HEADER = [name for name, _, _ in COLUMNS]
_KEYS = [key for _, _, key in COLUMNS]
_CONTENT_COL = _KEYS.index("content") + 1
_MEMBER_COL = _KEYS.index("member") + 1

# 국회의원 현황 시트에서 의원명(C열)을 기준으로 찾을 항목과 VLOOKUP 열 위치
# (C=의원명, D=정당, E=소속위원회, F=지역)
_LOOKUP_OFFSET = {"party": 2, "committee": 3, "district": 4}

# 프론트 결과 미리보기가 사용할 (키, 라벨) 목록 — COLUMNS가 단일 출처
UI_FIELDS = [
    {"key": key, "label": name} for name, _, key in COLUMNS if key and key != "content"
]
CONTENT_LABEL = HEADER[_CONTENT_COL - 1]

# openpyxl은 무거우므로 첫 사용 시점에 로딩한다 (서버 시작 속도 개선)
_STYLES: dict = {}
_MEMBERS: list[list[str]] = []  # [0]=헤더, [1:]=의원 행
_MEMBER_INFO: dict[str, tuple[str, str, str]] = {}  # 의원명 -> (정당, 소속위원회, 지역)


def _members() -> list[list[str]]:
    if not _MEMBERS:
        with open(_MEMBER_CSV, newline="", encoding="utf-8") as f:
            _MEMBERS.extend(csv.reader(f))
    return _MEMBERS


def member_info(name: str) -> tuple[str, str, str] | None:
    """의원명으로 (정당, 소속위원회, 지역)을 찾는다. 없으면 None."""
    if not _MEMBER_INFO:
        for row in _members()[1:]:
            _MEMBER_INFO.setdefault(row[2], (row[3], row[4], row[5]))
    return _MEMBER_INFO.get(name)


def enrich_member_fields(data: dict) -> dict:
    """요구의원이 국회의원 현황에 있으면 정당·소관위원회·지역구를 명단 값으로 채운다."""
    info = member_info(data.get("member", ""))
    if info:
        data["party"], data["committee"], data["district"] = info
    return data


def _styles() -> dict:
    if not _STYLES:
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

        side = Side(style="thin", color="CCCCCC")
        _STYLES.update(
            header_fill=PatternFill("solid", fgColor="1F6FB2"),
            header_font=Font(bold=True, color="FFFFFF"),
            border=Border(left=side, right=side, top=side, bottom=side),
            align_header=Alignment(horizontal="center", vertical="center"),
            align_plain=Alignment(vertical="center"),
            align_wrap=Alignment(vertical="center", wrap_text=True),
        )
    return _STYLES


def _sheet(wb):
    return wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active


class LedgerFormatError(ValueError):
    """기존 엑셀 파일의 열 구조가 관리대장 서식과 다를 때 발생."""


def _add_member_sheet(wb):
    from openpyxl.utils import get_column_letter

    st = _styles()
    ws = wb.create_sheet(MEMBER_SHEET_NAME)
    widths = [6, 8, 10, 14, 34, 20, 6, 8, 10]
    for r, row in enumerate(_members(), start=1):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            if r == 1:
                cell.fill = st["header_fill"]
                cell.font = st["header_font"]
                cell.alignment = st["align_header"]
            cell.border = st["border"]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    return ws


def create_ledger():
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    st = _styles()
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    for idx, (name, width, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=idx, value=name)
        cell.fill = st["header_fill"]
        cell.font = st["header_font"]
        cell.alignment = st["align_header"]
        cell.border = st["border"]
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    _add_member_sheet(wb)
    return wb


def open_ledger(path: str):
    from openpyxl import load_workbook

    wb = load_workbook(path)
    ws = _sheet(wb)
    header = [
        (str(ws.cell(row=1, column=i + 1).value).strip() if ws.cell(row=1, column=i + 1).value else "")
        for i in range(len(HEADER))
    ]
    if header != HEADER:
        raise LedgerFormatError(
            "선택한 엑셀 파일이 관리대장 서식과 다릅니다. "
            "'새 관리대장 생성'으로 다시 만들거나 올바른 파일을 선택해 주세요."
        )
    # 예전 버전으로 만든 파일에는 의원 명단 시트가 없을 수 있으므로 보충한다
    if MEMBER_SHEET_NAME not in wb.sheetnames:
        _add_member_sheet(wb)
    return wb


def _lookup_formula(key: str, row_idx: int, fallback: str) -> str:
    """의원명 셀을 키로 한 VLOOKUP 수식. 명단에 없으면 파싱한 값(fallback)을 쓴다."""
    from openpyxl.utils import get_column_letter

    member_cell = f"${get_column_letter(_MEMBER_COL)}{row_idx}"
    rng = f"'{MEMBER_SHEET_NAME}'!$C$2:$F${len(_members())}"
    fb = str(fallback).replace('"', '""')
    return f'=IFERROR(VLOOKUP({member_cell},{rng},{_LOOKUP_OFFSET[key]},0),"{fb}")'


def append_rows(wb, rows: list[dict]) -> int:
    """요구서 1건당 1행씩 관리대장에 추가하고, 추가된 행 수를 반환한다."""
    st = _styles()
    ws = _sheet(wb)
    next_row = ws.max_row + 1

    for offset, row_data in enumerate(rows):
        row_idx = next_row + offset
        for col, (_, _, key) in enumerate(COLUMNS, start=1):
            value = row_data.get(key, "") if key else ""
            if key in _LOOKUP_OFFSET:
                value = _lookup_formula(key, row_idx, value)
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.border = st["border"]
            cell.alignment = st["align_wrap"] if col == _CONTENT_COL else st["align_plain"]
    return len(rows)
